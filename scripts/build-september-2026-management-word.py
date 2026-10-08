"""Versión ejecutiva del informe de septiembre con tipografía legible."""
from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
OUT = ART / "informe-justificacion-gastos-septiembre-2026.docx"


def data(name):
    return json.loads((ART / name).read_text())


details = data("september-2026-details.json")
scenes = {x["sessionId"]: x for x in data("september-2026-scenes.json")["sessions"]}
media = {x["id"]: x for x in data("september-2026-podcaster-media.json")["rows"]}
video_times = {x["id"]: x for x in data("september-2026-video-times.json")["rows"]}
podcast = [x for x in details["podcaster"]["sessions"] if not x["id"].startswith("schroeder_")]
test_session_ids = {"hLYPCxeqPj0m0wJkakNS", "v8Kl45ENB4xhaPZdN08P"}
rate_veo = 8387 / 2496
rate_image = 3718.05 / 770
timezone = ZoneInfo("America/Cancun")


def local_day(timestamp):
    return datetime.fromisoformat(timestamp.replace("Z", "+00:00")).astimezone(timezone).date().isoformat()


def cash(n):
    return f"${n:,.2f}"


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.7)
sec.left_margin = Cm(2.15)
sec.right_margin = Cm(2.15)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(15 if name == "Title" else 13.5 if name in ("Subtitle", "Heading 1", "Heading 2") else 12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if name == "Subtitle":
        style.font.italic = False
    style.paragraph_format.space_after = Pt(7 if name == "Normal" else 6)
    if name not in ("Normal", "Subtitle"):
        style.font.bold = True
        style.paragraph_format.space_before = Pt(11 if name != "Title" else 0)
    if name == "Title":
        ppr = style._element.get_or_add_pPr()
        for child in list(ppr):
            if child.tag == qn("w:pBdr"):
                ppr.remove(child)


def para(text):
    parts = text.split("**")
    if len(parts) % 2 == 0:
        raise ValueError("Las negritas del párrafo deben estar cerradas")
    paragraph = doc.add_paragraph()
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        run.bold = index % 2 == 1
    return paragraph


def heading(text, level=1):
    doc.add_heading(text, level=level)


def shade(cell):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    cell._tc.get_or_add_tcPr().append(shd)


def pad(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, twips in (("top", 65), ("bottom", 65), ("left", 115), ("right", 115)):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(twips))
        margin.set(qn("w:type"), "dxa")
        tc_mar.append(margin)


def link(cell, href):
    p = cell.paragraphs[0]
    rid = p.part.relate_to(href, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    rpr.append(size)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "Abrir"
    r.append(t)
    h.append(r)
    p._p.append(h)


def inline_link(paragraph, label, href):
    rid = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    rpr.append(size)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = label
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for i, title in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = title
        shade(cell)
        pad(cell)
        cell.paragraphs[0].paragraph_format.keep_with_next = True
        if widths:
            cell.width = Cm(widths[i])
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0, 0, 0)
    header_repeat = OxmlElement("w:tblHeader")
    header_repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
    for values in rows:
        row = table.add_row()
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for i, value in enumerate(values):
            cell = row.cells[i]
            pad(cell)
            if widths:
                cell.width = Cm(widths[i])
            if isinstance(value, tuple):
                link(cell, value[0])
            else:
                cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(11)
                    run.font.color.rgb = RGBColor(0, 0, 0)
    return table


doc.add_paragraph("Informe de productividad y gastos de septiembre de 2026", style="Title")
doc.add_paragraph("CharlyBrown | Solo actividad y gastos de septiembre de 2026 | Corte: 30 de septiembre, 11:29 h de Cancún | Pesos mexicanos", style="Subtitle")
para("Tras varios meses de construir y ajustar las herramientas, **dos personas** crearon y mejoraron mucho más contenido en septiembre. Ese aumento del trabajo de producción elevó el uso de los servicios de inteligencia artificial y explica la mayor parte del gasto del mes.")

heading("1.- Producción y avance")
para("**Las mismas dos personas** trabajaron en videos y escape rooms: eligieron temas, prepararon materiales y revisaron resultados. Cada video de escena guardado es una generación o una nueva versión, no una sesión completa. PigPen había permitido crear 67 temas en julio; tras las pruebas de agosto, la producción creció en septiembre.")
add_table(["Medida", "Agosto", "Septiembre", "Cambio"], [
    ["Videos de escena guardados", "187", "312", "+66.8%"],
    ["Sesiones nuevas de video", "13", "26", "+100%"],
    ["Temas de escape room creados", "5", "74", "69 temas más"],
], [6.5, 3.2, 3.2, 3.5])

para("**Videos y audio.** Snoopy Studio (Podcaster) organizó 26 sesiones con 285 escenas. Ya hay video para 237 escenas y se guardaron **75 versiones adicionales** para corregirlas o mejorarlas. El trabajo incluyó **457 grabaciones de voz**, 241 imágenes de referencia y 115 modificaciones de esas imágenes. Schroeder Sound Lab, implementado en septiembre, prepara música de fondo original para los videos.")
para("**Escape rooms.** PigPen permitió crear 266 misiones, 1,064 preguntas y **355 elementos visuales**. En Inglés, el trabajo de septiembre cubrió el trimestre 1 y comenzó el trimestre 2. Los escape rooms de las asignaturas impartidas en español corresponden al trimestre 2.")
para("**Otros materiales.** Image Creator generó 36 imágenes; Marcie Blog Editor, 13 versiones de artículos de prueba; y Charly MCPEditor, **5 unidades de libro de prueba**. Una de esas unidades ya reúne 16 actividades, 20 recursos, 16 notas docentes, una lectura y 23 imágenes. También se registraron 2 lecturas independientes, 2 cursos Moodle y 2 pruebas del mismo simulador.")
para("**El equipo editorial aún no ha revisado el material.** Los cambios que solicite después, en especial nuevas versiones de escenas, generarán costo adicional.")

add_table(["Trabajo planeado", "Previsto", "Con avance", "Avance"], [
    ["Sesiones de video de Secundaria, trimestre 2", "63", "22 vinculadas a la planeación", "34.9%"],
    ["Escape rooms de asignaturas en español", "65", "39 temas con misiones", "60.0%"],
    ["Escape rooms de Inglés planeados para trimestres 2 y 3", "30", "11 capítulos con misiones", "36.7%"],
    ["Actividades científicas y simuladores", "34", "Aún en preparación", "0%"],
], [7.2, 2.2, 4.5, 2.5])
para("La hoja general abarca varias asignaturas impartidas en español. En Inglés, esta tabla mide solo el avance del plan de trimestres 2 y 3; el trabajo del trimestre 1 se cuenta en la producción de septiembre, pero no en ese porcentaje.")

heading("2.- Gastos de septiembre")
para("La mayoría del gasto corresponde al **trabajo de creación**: cada escena, voz, imagen, texto o búsqueda utiliza un servicio que se cobra por uso. El funcionamiento de las herramientas creadas fue cerca del **6%** de Google Cloud.")
gastos = add_table(["Grupo de gasto", "Importe", "Para qué se usó"], [
    ["Creación con inteligencia artificial", "$19,090.89", "Video, imágenes, texto, búsquedas y audio"],
    ["Funcionamiento de las herramientas", "$1,186.44", "Ejecución, publicación y almacenamiento"],
    ["Total de Google Cloud", "$20,277.33", "Total registrado hasta el corte"],
], [6.0, 3.0, 7.4])
for row in (gastos.rows[1], gastos.rows[3]):
    for cell in row.cells[:2]:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
para("El gasto de creación incluye $8,441.28 en video, $3,921.77 en búsquedas con Google Grounding, $3,718.05 en imágenes y $2,936.65 en texto; el resto corresponde a audio y otros usos menores. **Marcie Blog Editor usó Grounding en pruebas de artículos**, pero esas búsquedas ya se retiraron para reducir costos. La factura no separa cuánto de los $3,921.77 corresponde a Marcie o a cada artículo.")
para("De los $8,441.28 de video, $8,387.00 corresponden a piezas con duración registrada: $5,430.04 estimados para rmora@asc.education y $2,956.96 para wlopez@asc.education, según los segundos guardados. Los otros $54.28 no se pudieron repartir así. La factura tampoco asigna el gasto de textos e imágenes a cada persona o escape room.")

heading("3.- Detalle de la producción")
heading("Producción diaria de videos", 2)
para("Una sesión se cuenta como completa cuando **todas sus escenas tienen video y voz guardados**. Esta medida aún no incluye la revisión editorial. El conteo de sesiones nuevas excluye dos pruebas vacías.")
new_sessions_by_day = Counter(local_day(x["createdAt"]) for x in podcast if x["id"] not in test_session_ids)
worked_sessions_by_day = Counter(day for x in video_times.values() for day in x["activeVideoDays"])
complete_by_day = Counter()
complete_same_day = Counter()
complete_sessions = []
for session in podcast:
    scene = scenes.get(session["id"], {})
    voice = media.get(session["id"], {})
    timing = video_times.get(session["id"], {})
    scene_count = session["scenes"]
    if not scene_count or scene.get("distinctReadyScenes", 0) != scene_count or voice.get("distinctRowsWithReadyAudio", 0) != scene_count:
        continue
    if not timing.get("lastVideoAt") or not timing.get("lastAudioAt"):
        continue
    finished_day = local_day(max(timing["lastVideoAt"], timing["lastAudioAt"]))
    complete_by_day[finished_day] += 1
    if finished_day == local_day(timing["createdAt"]):
        complete_same_day[finished_day] += 1
    complete_sessions.append((session["id"], scene_count, finished_day))
assert len(complete_sessions) == 2 and sum(complete_same_day.values()) == 1, "Revisar sesiones completas y fechas"
daily_rows = []
for day in sorted(set(new_sessions_by_day) | set(worked_sessions_by_day) | set(complete_by_day)):
    finished = [count for _id, count, finished_day in complete_sessions if finished_day == day]
    ready_label = ", ".join(f"1 ({count} escena" + ("s" if count != 1 else "") + ")" for count in finished) if finished else "0"
    daily_rows.append([day[8:10] + "/09", new_sessions_by_day[day], worked_sessions_by_day[day], ready_label])
add_table(["Día", "Sesiones nuevas", "Sesiones con video guardado", "Sesiones completas con video y voz"], daily_rows, [2.2, 3.5, 5.0, 5.7])
para("**Capacidad observada:** el día de mayor actividad se crearon 4 sesiones y se trabajó en video de hasta 4 sesiones. Solo **una sesión de una escena** quedó completa el mismo día de su creación. La sesión completa de **14 escenas** se inició el 29 y terminó el 30 de septiembre. Estos registros muestran la producción real del equipo, pero aún no permiten afirmar cuántas sesiones largas puede terminar la herramienta en un solo día de forma constante.")
heading("Sesiones de video", 2)
para("Cada fila indica escenas con video, videos y voces guardados, costo aproximado y enlace para compartir la sesión. El costo suma video e imágenes de referencia; algunas sesiones solo tienen gasto de imágenes. Las dos sesiones de prueba se agrupan sin enlace.")
video_rows = []
test_sessions = [x for x in podcast if x["id"] in test_session_ids]
assert len(test_sessions) == 2, "Deben existir las dos sesiones de prueba"
for x in podcast:
    if x["id"] in test_session_ids:
        continue
    s = scenes.get(x["id"], {})
    m = media.get(x["id"], {})
    cost = s.get("readyVideos", 0) * 8 * rate_veo + (m.get("rowReferenceImageListItems", 0) + m.get("recordedReferenceEdits", 0)) * rate_image
    video_rows.append([
        f"{x['title']} ({x['user'].split('@')[0]})",
        f"{s.get('distinctReadyScenes', 0)}/{x['scenes']} escenas con video; {s.get('readyVideos', 0)} videos; {m.get('readyAudioJobs', 0)} voces. Costo aprox.: {cash(cost)}",
        (f"https://charly-brown.web.app/video-player.html?sessionId={x['id']}",),
    ])
test_scenes = sum(x["scenes"] for x in test_sessions)
test_videos = sum(scenes.get(x["id"], {}).get("readyVideos", 0) for x in test_sessions)
test_ready_scenes = sum(scenes.get(x["id"], {}).get("distinctReadyScenes", 0) for x in test_sessions)
test_voices = sum(media.get(x["id"], {}).get("readyAudioJobs", 0) for x in test_sessions)
test_cost = sum(
    round(
        scenes.get(x["id"], {}).get("readyVideos", 0) * 8 * rate_veo
        + (media.get(x["id"], {}).get("rowReferenceImageListItems", 0)
           + media.get(x["id"], {}).get("recordedReferenceEdits", 0)) * rate_image,
        2,
    )
    for x in test_sessions
)
video_rows.append([
    "Sesiones de prueba (wlopez)",
    f"2 sesiones; {test_ready_scenes}/{test_scenes} escenas con video; {test_videos} videos; {test_voices} voces. Costo aprox.: {cash(test_cost)}",
    "",
])
add_table(["Sesión y creador", "Trabajo guardado", "Enlace"], video_rows, [5.8, 8.3, 2.3])

heading("Temas de escape room", 2)
para("Cada tema muestra el material guardado por PigPen y un **costo parcial de elementos visuales**, calculado con el costo promedio de imagen del proyecto. La factura no separa por tema el gasto de texto para objetivos, misiones y preguntas. La lista incluye versiones repetidas; por eso sus 74 registros no equivalen a 74 temas curriculares únicos.")
topic_rows = []
topics = sorted(details["pigpen"]["topics"], key=lambda x: (
    0 if x["materia"] == "Inglés" else 1,
    x["materia"], str(x["trimestre"]), x["grado"], str(x["tema"]), x["title"],
))
for x in topics:
    visual_count = x["missionImages"] + int(x["coverImage"]) + int(x["endingImage"])
    topic_rows.append([
        f"{x['title']} ({x['user'].split('@')[0]})",
        f"{x['materia'] or 'Borrador'}; {x['grado']}; T{x['trimestre']}; tema {x['tema']}",
        f"{x['missions']} misiones; {x['questions']} preguntas; {visual_count} visuales",
        cash(visual_count * rate_image),
    ])
total_visuals = sum(x["missionImages"] + int(x["coverImage"]) + int(x["endingImage"]) for x in topics)
for user in ("rmora@asc.education", "wlopez@asc.education"):
    own = [x for x in topics if x["user"] == user]
    own_visuals = sum(x["missionImages"] + int(x["coverImage"]) + int(x["endingImage"]) for x in own)
    own_cost = sum(round((x["missionImages"] + int(x["coverImage"]) + int(x["endingImage"])) * rate_image, 2) for x in own)
    topic_rows.append([f"Subtotal {user.split('@')[0]}", f"{len(own)} versiones", f"{sum(x['missions'] for x in own)} misiones; {own_visuals} visuales", f"≈ {cash(own_cost)}"])
total_cost = sum(round((x["missionImages"] + int(x["coverImage"]) + int(x["endingImage"])) * rate_image, 2) for x in topics)
topic_rows.append(["Total de 74 versiones", "Inglés y asignaturas en español", "266 misiones; 1,064 preguntas; 355 visuales", f"≈ {cash(total_cost)}"])
add_table(["Tema y creador", "Asignatura y curso", "Material guardado", "Costo visual"], topic_rows, [5.4, 4.0, 4.4, 2.6])

heading("4.- Costo por producto")
para("Los costos por pieza son parciales: los $2,936.65 de texto no se pueden repartir por tema.")
add_table(["Producto", "Resultado de septiembre", "Costo parcial por pieza"], [
    ["Video", "1 sesión con 14/14 escenas en video y voz; 26 sesiones y 312 videos de escena", "$605.23 en video e imágenes de apoyo"],
    ["Escape room de Inglés", "26 armados (T1 y T2); 29 versiones con misiones", "$24.70 en imágenes; texto sin desglose"],
    ["Escape room de asignaturas en español", "10 armados; 44 versiones con misiones", "$28.49 en imágenes; texto sin desglose"],
    ["Artículo de Marcie", "13 versiones de prueba; revisión pendiente", "Grounding generó gasto; sin desglose por artículo"],
    ["Unidad de libro", "5 pruebas; 1 con actividades, recursos, notas y lectura", "$111.06 en imágenes de esa unidad"],
], [4.0, 6.7, 5.7])
heading("Proyección para completar el trimestre 2", 2)
para("La planeación permite identificar **41 guiones de video sin sesión enlazada**, **26 temas de escape room en español sin misiones** y **4 temas de Inglés del trimestre 2 sin misiones**. Para estimar el gasto se aplica a cada pendiente el costo parcial observado en septiembre para un producto comparable. En video, el ejemplo es una sesión de 14 escenas con video y voz; los guiones pendientes pueden requerir otra cantidad de escenas.")
add_table(["Trabajo pendiente", "Piezas", "Referencia parcial por pieza", "Proyección parcial"], [
    ["Videos del trimestre 2", "41", "$605.23 por sesión de 14 escenas", "$24,814.43"],
    ["Escape rooms en español, trimestre 2", "26", "$28.49 en imágenes", "$740.74"],
    ["Escape rooms de Inglés, trimestre 2", "4", "$24.70 en imágenes", "$98.80"],
    ["Total indicativo", "71 piezas", "Video e imágenes de apoyo", "$25,653.97"],
], [5.0, 2.0, 5.2, 4.2])
para("**Esta cifra no es un presupuesto completo ni un importe ya gastado.** Faltan costos de voz, textos, objetivos, misiones, preguntas y otros usos que la factura no separa por pieza. Tampoco incluye terminar las 22 sesiones de video ya enlazadas que sigan incompletas. La revisión editorial podrá pedir cambios: cada nueva generación de una escena de video de ocho segundos costaría aproximadamente **$26.88** con la tarifa observada, más los servicios de voz o imagen que se vuelvan a utilizar. Por eso el gasto final puede variar.")
plan_source = doc.add_paragraph("Fuente de las piezas pendientes: ")
inline_link(plan_source, "planeación de videos y escape rooms", details.get("planSource", "https://docs.google.com/spreadsheets/d/1KeIukb-Cu_iv9eiJii3Jg1O2P4-bMgaE7K2OvKSgerg/edit"))
plan_source.add_run(" y registros de producción de septiembre.")
heading("Escenario preliminar para el trimestre 3", 2)
para("**El trimestre 3 todavía no tiene un plan integral aprobado.** Para dimensionar el posible gasto, este escenario supone provisionalmente el mismo volumen que el plan del trimestre 2: **63 sesiones de video** y **65 escape rooms de asignaturas en español**. La hoja de Inglés ya enumera **15 capítulos** para el trimestre 3. Estas cantidades son una referencia para presupuestar, no metas ni trabajos autorizados.")
t3_video_cost = 63 * 605.23
t3_spanish_cost = 65 * 28.49
t3_english_cost = 15 * 24.70
t3_video_images = 63 * 14
t3_spanish_images = round(t3_spanish_cost / rate_image)
t3_english_images = round(t3_english_cost / rate_image)
add_table(["Producto hipotético", "Cantidad", "Imágenes estimadas", "Gasto parcial estimado"], [
    ["Sesiones de video", "63", f"{t3_video_images:,} de referencia", cash(t3_video_cost)],
    ["Escape rooms en español", "65", f"≈ {t3_spanish_images} elementos visuales", cash(t3_spanish_cost)],
    ["Escape rooms de Inglés", "15", f"≈ {t3_english_images} elementos visuales", cash(t3_english_cost)],
    ["Total del escenario", "143 productos", f"≈ {t3_video_images + t3_spanish_images + t3_english_images:,} imágenes y visuales", cash(t3_video_cost + t3_spanish_cost + t3_english_cost)],
], [4.4, 2.6, 4.7, 4.7])
para("De ese total, aproximadamente **$6,481.21 corresponderían a imágenes y elementos visuales**, ya incluidos en las cifras de la tabla; no se suman de nuevo. La proyección supone 14 imágenes de referencia por video y usa los costos parciales observados por producto en septiembre. No incluye voz, textos, nuevas versiones, cambios editoriales ni imágenes de libros, artículos o Image Creator, porque el trimestre 3 todavía no define cuántas piezas de esos tipos se harán. **La cifra se recalculará cuando exista la planeación del trimestre 3**, con escenas y temas reales.")
doc.add_page_break()
heading("Conclusión", 2)
para("**Septiembre demuestra el valor de la suite CharlyBrown en manos del equipo.** Tras meses de desarrollo, dos personas transformaron guiones y planeaciones propias en escenas con voz e imágenes, escape rooms con misiones y preguntas, y materiales para libros y cursos. Snoopy Studio, PigPen y las demás herramientas forman un sistema reutilizable: **cada nueva producción suma activos educativos para ASC**. El mayor uso de inteligencia artificial refleja más trabajo creativo y más contenido preparado para revisión editorial. El programa ya tiene una ruta concreta: avanzar en los videos y escape rooms pendientes del trimestre 2 y después planear los videos y escape rooms de las asignaturas en español para el trimestre 3.")
para("Los videos y escape rooms ya incorporan la identidad y el logotipo de **ASC**. Tras la revisión, los videos quedarán en **Vimeo** para la empresa y los escape rooms en **Moodle para México** y la versión de **Panamá prevista para inicios de 2027**. Estos materiales pueden enriquecer los cursos de ambos mercados y mejorar la experiencia de sus alumnos. Los guiones, la estructura de las sesiones y la automatización son aportaciones del equipo; Google no reclama la propiedad del contenido generado y la empresa podrá comercializarlo conforme a sus derechos y acuerdos aplicables.")
para("**Para contener el gasto, ASC puede priorizar uno o dos videos de cada una de las siete asignaturas con guiones pendientes**, en lugar de producir de inmediato los 41 videos restantes del trimestre 2. Serían **7 a 14 sesiones**, con un costo parcial estimado de **$4,236.61 a $8,473.22**, frente a **$24,814.43** para las 41; la diferencia parcial sería de **$16,341.21 a $20,577.82**. Esta opción permite probar qué videos aportan más valor antes de ampliar la producción. También implica menos material audiovisual para apoyar y ampliar los temas, lo que podría afectar la experiencia y la percepción de docentes y alumnos. Conviene elegir los temas prioritarios con el equipo editorial y valorar su uso antes de decidir sobre los demás.")
para("**El trimestre 3 aún no cuenta con una planeación integral.** Aunque la hoja de Inglés enumera capítulos para ese trimestre, faltan los planes de videos y escape rooms de las asignaturas en español. Cuando esté definido su alcance, la producción real de los trimestres 1 y 2 permitirá calcular una proyección más sólida de los gastos futuros. Así, ASC podrá decidir cuánto producir y revisar con una base cada vez más precisa, aprovechando una suite que ya mostró capacidad para convertir el trabajo de un equipo pequeño en materiales propios, útiles y preparados para llegar a más alumnos.")
sources = doc.add_paragraph("Fuentes: ")
inline_link(sources, "Términos de Gemini API", "https://ai.google.dev/gemini-api/terms")
sources.add_run(" y ")
inline_link(sources, "Google Cloud", "https://cloud.google.com/terms/service-terms")

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
