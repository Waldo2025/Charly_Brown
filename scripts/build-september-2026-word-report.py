"""Informe sencillo para dirección a partir de los inventarios de septiembre."""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
ART = ROOT / "artifacts"
OUT = ART / "informe-justificacion-gastos-septiembre-2026.docx"


def load(name):
    return json.loads((ART / name).read_text())


details = load("september-2026-details.json")
scenes = load("september-2026-scenes.json")
media = load("september-2026-podcaster-media.json")
charly = load("september-2026-charly-details.json")
reading = load("september-2026-reading.json")
by_scene = {x["sessionId"]: x for x in scenes["sessions"]}
by_media = {x["id"]: x for x in media["rows"]}
podcast = [x for x in details["podcaster"]["sessions"] if not x["id"].startswith("schroeder_")]
rate_video = 8387 / 2496
rate_image = 3718.05 / 770


def mxn(value):
    return f"${value:,.2f}"


def date(value):
    from datetime import datetime, timezone, timedelta
    return datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone(timezone(timedelta(hours=-5))).strftime("%d/%m/%Y")


def visual_count(x):
    return x.get("missionImages", 0) + int(bool(x.get("coverImage"))) + int(bool(x.get("endingImage")))


def charly_images(unit):
    return sum(v for k, v in unit.get("assetKinds", {}).items() if k.startswith("image/"))


doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(2.1)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.1)
sec.right_margin = Cm(2.1)

styles = doc.styles
styles["Normal"].font.name = "DejaVu Sans"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].paragraph_format.space_after = Pt(6)
for style_name, size, before, after in [("Title", 17, 0, 12), ("Heading 1", 13, 12, 6), ("Heading 2", 11.5, 9, 4)]:
    st = styles[style_name]
    st.font.name = "DejaVu Sans"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    if style_name == "Title":
        ppr = st._element.get_or_add_pPr()
        for child in list(ppr):
            if child.tag == qn("w:pBdr"):
                ppr.remove(child)


def paragraph(text="", bold_start=None):
    p = doc.add_paragraph()
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold = True
        p.add_run(text[len(bold_start):])
    else:
        p.add_run(text)
    return p


def heading(text, level=1):
    return doc.add_heading(text, level)


def shade(cell, fill="F0F0F0"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def hyperlink(cell, label, url):
    p = cell.paragraphs[0]
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "165A86")
    rpr.append(color)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = label
    run.append(t)
    h.append(run)
    p._p.append(h)


def table(headers, rows, small=False):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, name in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = str(name)
        shade(c)
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = "DejaVu Sans"
                r.font.size = Pt(8.6 if small else 9.2)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        new_row = t.add_row()
        cant_split = OxmlElement("w:cantSplit")
        new_row._tr.get_or_add_trPr().append(cant_split)
        cells = new_row.cells
        for i, value in enumerate(row):
            if isinstance(value, tuple) and len(value) == 2 and value[0].startswith("https://"):
                hyperlink(cells[i], value[1], value[0])
            else:
                cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "DejaVu Sans"
                    run.font.size = Pt(8.2 if small else 9.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


title = doc.add_paragraph(style="Title")
title.add_run("Justificación de producción y gastos de septiembre de 2026")
paragraph("Proyecto CharlyBrown | Corte de actividad: 30 de septiembre, 11:29 h de Cancún | Importes en pesos mexicanos")
paragraph("Este informe explica qué produjo el equipo durante septiembre, cómo aumentó la producción frente a los meses de prueba y qué parte del gasto de Google Cloud se puede relacionar con los materiales guardados. Se preparó para facilitar la revisión de dirección. Los importes del mes todavía son provisionales.")

heading("1 Resumen para dirección")
paragraph("Google Cloud registró $20,277.33 durante septiembre. Vertex AI, que incluye los modelos de video, imagen y texto, representó $19,090.89, alrededor del 94% del total. Los servicios que mantienen la herramienta funcionando sumaron $1,186.44, alrededor del 6%. El gasto creció principalmente porque ya se usaron las herramientas para producir materiales, después de varios meses de desarrollo y pruebas.")
paragraph("Dos personas crearon las 26 sesiones de video de Podcaster y dos personas crearon los 74 temas de escape room de PigPen. En video se guardaron 312 generaciones de escenas en septiembre, frente a 187 en agosto: un aumento del 66.8%. Las sesiones nuevas de Podcaster pasaron de 13 a 26, el doble. PigPen pasó de 5 temas creados en agosto a 74 en septiembre; julio ya había tenido 67, por lo que esa comparación se presenta como referencia y no como una subida constante.")
table(["Producto", "Trabajo guardado en septiembre", "Responsables"], [
    ["Video", "26 sesiones, 285 escenas registradas, 312 videos de escena listos", "2 creadores"],
    ["Escape rooms", "74 temas, 266 misiones, 1,064 preguntas y 355 espacios con imagen", "2 creadores"],
    ["Imagen independiente", "36 imágenes en 11 sesiones de Image Creator", "3 creadores"],
    ["Artículos", "13 versiones con contenido en 8 sesiones de Marcie; son pruebas", "3 creadores"],
    ["Unidades de libro", "5 unidades guardadas en Charly MCPEditor; 1 con amplio contenido", "3 creadores"],
    ["Otros", "2 cursos Moodle, 2 lecturas y 2 registros de prueba de simulador", "Varios creadores"],
])

heading("2 Qué se produjo y por qué se usó la inteligencia artificial")
heading("Videos y sonido", 2)
paragraph("Cada video generado por Podcaster corresponde a una escena, no a una sesión completa. Las 26 sesiones tienen 285 escenas registradas; 237 escenas distintas recibieron al menos un video y otras 75 generaciones se usaron para probar o sustituir resultados. También se guardaron 457 voces en off, incluidas 176 generaciones adicionales, y 241 imágenes de referencia. El historial de esas referencias muestra 115 modificaciones. Este trabajo permite montar videos educativos con guion, imagen y audio coordinados.")
paragraph("Una sesión de 14 escenas tiene video y voz en todas sus escenas. El resto de las sesiones necesita verificar las escenas que no tienen video antes de presentarlas como videos completos. No se observó en los datos una aprobación editorial o un archivo final exportado; por eso 'completo' en este informe significa cobertura técnica de escenas, no publicación ya autorizada.")
paragraph("Schroeder Sound Lab se usa para crear música de fondo. Se identificó una sesión propia de música y 10 trabajos musicales listos en Podcaster; no hay datos para afirmar que sean 10 pistas distintas creadas por Schroeder ni para asignarles un costo individual.")

heading("Escape rooms y material visual", 2)
paragraph("PigPen guardó 74 temas: 29 de Inglés, 44 de asignaturas impartidas en español y un borrador sin asignatura. En conjunto conservan 266 misiones, 1,064 preguntas y 355 espacios con imagen: 247 imágenes de misión, 71 portadas y 37 cierres. Se consideran equipados para esta comparación los temas con al menos 4 misiones, 16 preguntas, 4 imágenes de misión y portada: 26 de Inglés y 10 de asignaturas en español. Los demás requieren una revisión individual, ya que algunos podrían estar diseñados con menos misiones.")
paragraph("Las imágenes cumplen distintas funciones: PigPen ilustra pistas y escenarios; Charly MCPEditor incorpora recursos en unidades de libro; Podcaster prepara referencias para las escenas; Image Creator produce imágenes independientes; Science Activities servirá para actividades y simuladores. No todo espacio con imagen equivale a una llamada cobrada por la IA: algunas referencias se pueden importar o reutilizar.")

heading("Artículos, unidades, lecturas y cursos", 2)
paragraph("Marcie Blog Editor conserva 13 versiones de artículo con contenido repartidas en 8 sesiones nuevas. Para este informe se consideran pruebas, aunque dos tengan un estado interno de publicación. Charly MCPEditor tiene 5 unidades guardadas. Una unidad, 'Un proyecto interesante', conserva 16 actividades, 20 recursos, 16 notas para docentes, una lectura y 23 archivos de imagen; seis archivos adicionales son PDF. Las otras unidades tienen distintos grados de avance. También se guardaron 2 lecturas y 2 cursos Moodle.")
paragraph("Science Activities conserva dos registros de prueba del mismo simulador. Ninguno corresponde todavía a los 26 trabajos de actividades ni a los 8 simuladores previstos para el segundo trimestre. Por eso el avance de ese plan se mantiene en 0%.")

heading("3 Avance frente a la planeación")
table(["Trabajo previsto", "Plan", "Con evidencia", "Pendiente", "Avance"], [
    ["Guiones de video de Secundaria, trimestre 2", 63, "22 sesiones enlazadas", 41, "34.9%"],
    ["Escape rooms de asignaturas en español, trimestre 2", 65, "39 temas con misiones", 26, "60.0%"],
    ["Escape rooms de Inglés, trimestres 2 y 3", 30, "11 capítulos con misiones", 19, "36.7%"],
    ["Actividades científicas, trimestre 2", 26, 0, 26, "0%"],
    ["Simuladores, trimestre 2", 8, 0, 8, "0%"],
])
paragraph("La hoja general de escape rooms corresponde a asignaturas impartidas en español, entre ellas Español, Matemáticas, Geografía, Biología e Historia. La hoja de Inglés se cuenta por separado. La planeación y el inventario pueden diferir porque la hoja aún marca como pendientes algunos proyectos que ya tienen misiones guardadas.")

heading("4 Detalle del gasto")
table(["Servicio de Google Cloud", "Gasto de septiembre"], [
    ["Vertex AI: creación de video, imagen, texto y búsquedas", "$19,090.89"],
    ["Cloud Run y Cloud Run Functions: ejecución de procesos", "$1,001.44"],
    ["Firebase Hosting: publicación de las herramientas", "$67.63"],
    ["App Engine", "$60.16"],
    ["Voz de Google Cloud", "$32.10"],
    ["Artifact Registry, almacenamiento y programador", "$25.09"],
    ["Ajuste de redondeo entre renglones", "$0.02"],
    ["Total mostrado por Cloud Billing", "$20,277.33"],
])
table(["Uso dentro de Vertex AI", "Gasto"], [
    ["Video Veo, incluidos dos tipos de generación", "$8,441.28"],
    ["Búsquedas usadas por los modelos", "$3,921.77"],
    ["Imágenes generadas por Gemini", "$3,718.05"],
    ["Texto generado o procesado por Gemini", "$2,936.65"],
    ["Audio y otros usos menores", "$73.13"],
    ["Ajuste de redondeo", "$0.01"],
    ["Total Vertex AI", "$19,090.89"],
])
paragraph("La factura agrupa cargos por proyecto y tipo de servicio; no indica qué empleado ni qué editor hizo cada llamada. En el renglón principal de Veo ($8,387.00) sí es posible hacer un reparto aproximado por los segundos de video guardados: $5,430.04 para las sesiones de rmora@asc.education y $2,956.96 para las de wlopez@asc.education. Es una estimación de ese tipo de gasto, no un cobro individual exacto.")
paragraph("Para ilustrar el gasto visual, se repartieron los $3,718.05 de imágenes entre 770 imágenes o modificaciones observables. Ese reparto equivale a unos $4.83 por unidad visual. Puede incluir referencias importadas y no refleja diferencias de precio entre modelos. Las cifras por producto son parciales y orientativas. Los costos de texto, búsquedas, audio, infraestructura y el servicio Render no se pueden repartir con seguridad por creador. Render quedó pendiente de consultar por falta de acceso a su factura.")

heading("5 Inventario de sesiones de video")
paragraph("La columna 'Videos' cuenta generaciones de una escena. 'Extras' son nuevas generaciones de escenas que ya tenían un video. El enlace usa la opción Compartir sesión de Podcaster y puede pedir acceso al sistema.")
video_rows = []
for x in podcast:
    a = by_scene.get(x["id"], {})
    m = by_media.get(x["id"], {})
    cost = a.get("readyVideos", 0) * 8 * rate_video + (m.get("rowReferenceImageListItems", 0) + m.get("recordedReferenceEdits", 0)) * rate_image
    video_rows.append([x["title"], x["user"].split("@")[0], f"{a.get('distinctReadyScenes', 0)}/{x['scenes']}", f"{a.get('readyVideos', 0)} ({a.get('extraReadyGenerations', 0)} extra)", m.get("readyAudioJobs", 0), mxn(cost), (f"https://charly-brown.web.app/video-player.html?sessionId={x['id']}", "Abrir")])
table(["Sesión", "Creador", "Escenas con video", "Videos", "Voces", "Costo parcial", "Enlace"], video_rows, small=True)

heading("6 Inventario de escape rooms")
paragraph("'Español' agrupa aquí las asignaturas impartidas en español; la columna Asignatura muestra cuál es cada una. Las imágenes son la suma de imágenes de misión, portada y cierre. El costo es un reparto orientativo del gasto de imágenes, sin texto ni búsquedas.")
pig_rows = []
for x in details["pigpen"]["topics"]:
    images = visual_count(x)
    subject = x.get("materia") or "Borrador"
    context = f"{subject} / {x.get('grado') or '?'} / T{x.get('trimestre') or '?'}"
    pig_rows.append([x["title"], x["user"].split("@")[0], context, x.get("missions", 0), images, mxn(images * rate_image)])
table(["Tema", "Creador", "Asignatura grado y trimestre", "Misiones", "Imágenes", "Costo parcial"], pig_rows, small=True)

heading("7 Otros materiales guardados")
heading("Unidades de Charly MCPEditor", 2)
unit_rows = []
for x in charly["rows"]:
    for u in x["units"] or [None]:
        n = charly_images(u) if u else 0
        unit_rows.append([x["owner"].split("@")[0], f"{x['title']} / {u['title'] if u else 'Sin unidad'}", u.get("activities", 0) if u else 0, f"{u.get('resources', 0) if u else 0} / {u.get('teacherNotes', 0) if u else 0}", n, mxn(n * rate_image)])
table(["Creador", "Sesión y unidad", "Actividades", "Recursos y notas", "Imágenes", "Costo parcial"], unit_rows, small=True)
paragraph("La unidad 'Un proyecto interesante' cubre proyectos, artes, lectura, escritura, matemáticas y otras áreas. Los importes de la tabla solo reflejan imágenes; no incluyen el trabajo de texto.")

heading("Imágenes de Image Creator", 2)
table(["Creador", "Sesión", "Imágenes", "Modelo usado", "Costo parcial"], [[x["user"].split("@")[0], x["title"], x["images"], ", ".join(f"{k}: {v}" for k, v in x["models"].items()) or "Sin imagen", mxn(x["images"] * rate_image)] for x in details["imageCreator"]["sessions"]], small=True)

heading("Artículos de Marcie Blog Editor", 2)
marcie_rows = []
for x in details["marcie"]["sessions"]:
    for v in x["versions"] or [None]:
        audience = {"educators": "docentes", "students": "alumnos", "parents": "familias", "coordinators": "coordinadores"}.get(v["audience"], v["audience"]) if v else "Sin versión"
        marcie_rows.append([x["user"].split("@")[0], v["title"] if v else f"{x['title']} (sin contenido)", audience, v["blocks"] if v else 0, "Prueba", "No calculable"])
table(["Creador", "Versión de artículo", "Público", "Bloques", "Situación", "Costo"], marcie_rows, small=True)

heading("Lecturas de Generar Lectura", 2)
readings = [(x, collection) for collection in ("lecturasASC", "lecturasNuevas") for x in reading["collections"][collection]["septemberRows"]]
table(["Creador", "Lectura", "Nivel", "Grado", "Trimestre", "Costo"], [[x["user"].split("@")[0], x["title"], x["level"], x["grade"], x["trim"], "No calculable"] for x, _ in readings], small=True)

heading("8 Alcance y fuentes")
paragraph("Los registros de actividad provienen de la base de datos del proyecto CharlyBrown, consultada el 30 de septiembre a las 11:29 h de Cancún. Los gastos provienen de la vista de Cloud Billing de septiembre y pueden cambiar cuando cierre la facturación. Los productos eliminados antes de la consulta no se pueden recuperar de este inventario. La comparación con la planeación usa el archivo de Google Sheets 'Guión videos Secundaria Aprende 2026-2027'.")
paragraph("Según los términos de Google Cloud y Gemini API, Google no reclama la propiedad del contenido nuevo generado por sus servicios. Esto permite su uso comercial conforme a los términos aplicables y a los derechos sobre los insumos. Los términos no garantizan exclusividad absoluta, pues otros clientes podrían obtener contenido similar. Fuentes: cloud.google.com/legal/archive/terms/service-terms/index-20260729 y ai.google.dev/gemini-api/terms.")

final_heading = heading("9 Contenido completo y costo aproximado por pieza")
final_heading.paragraph_format.page_break_before = True
paragraph("Esta tabla aplica un criterio conservador de contenido técnicamente armado. Aún hace falta la revisión editorial antes de declarar una pieza aprobada para venta o publicación. Los importes cubren solo la parte de video o imagen que se pudo distribuir; no son el costo total de producir cada pieza.")
table(["Tipo de contenido", "Cantidad con criterio técnico completo", "Criterio usado", "Gasto aproximado por pieza, parcial"], [
    ["Video completo", "1 sesión", "Video y voz en sus 14 escenas", "$605.23 para esa sesión; no incluye todo el audio y texto"],
    ["Escape room de Inglés", "26 temas", "4 misiones, 16 preguntas, 4 imágenes y portada", "$24.70 por tema, promedio de imágenes"],
    ["Escape room de asignaturas en español", "10 temas", "El mismo criterio visual y de preguntas", "$28.49 por tema, promedio de imágenes"],
    ["Artículo", "0 aprobados; 13 versiones de prueba", "No hay aprobación editorial registrada", "No calculable por artículo con los datos actuales"],
    ["Unidad de libro", "1 unidad con contenido amplio; 0 aprobadas", "16 actividades, 20 recursos, 16 notas y lectura", "$111.06 en imágenes de esa unidad; texto no calculable"],
], small=True)
paragraph("Un costo de $0.00 en una tabla de imágenes significa que no se encontró una imagen a la cual aplicar el reparto; no significa que la sesión haya sido gratuita. Para calcular el costo completo por pieza, cada herramienta deberá registrar en cada llamada el creador, la sesión, el modelo y las unidades cobradas, y luego cruzarlas con la factura.")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("CharlyBrown  |  Informe de septiembre de 2026").font.size = Pt(8)

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(OUT)
